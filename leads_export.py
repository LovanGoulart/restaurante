"""Exportação administrativa de clientes e preferências, sem alterar o banco."""

from collections import Counter, defaultdict
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
import re
import textwrap
from zoneinfo import ZoneInfo

from flask import Blueprint, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy.orm import joinedload, selectinload

from customers import role
from models import Cliente, Pedido, PedidoOnline, db

bp = Blueprint("leads_export", __name__)
BRASIL = ZoneInfo("America/Sao_Paulo")
MOEDA = '"R$" #,##0.00'
INVALIDOS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def texto(value):
    return INVALIDOS.sub("", str(value or ""))[:32767]


def data_local(value):
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(BRASIL).date()


def reais(centavos):
    return Decimal(centavos) / Decimal(100)


def resumir_cliente(cliente, pedidos, hoje):
    """Contabiliza cada pedido uma vez e soma quantidades por nome do item."""
    pedidos = sorted(pedidos, key=lambda p: (p.criado, p.id), reverse=True)
    frequencia = Counter()
    valores = Counter()
    presencas = Counter()
    nomes = {}
    for pedido in pedidos:
        presentes = set()
        for item in pedido.itens:
            nome = " ".join(item.nome.split())
            chave = nome.casefold()
            nomes.setdefault(chave, nome)
            frequencia[chave] += item.quantidade
            valores[chave] += item.quantidade * item.preco
            presentes.add(chave)
        presencas.update(presentes)

    ranking = sorted(
        frequencia,
        key=lambda k: (-frequencia[k], -presencas[k], k),
    )
    top = ranking[:3]
    ultima = data_local(pedidos[0].criado) if pedidos else None
    dias = max(0, (hoje - ultima).days) if ultima else None
    total = sum(p.total for p in pedidos)
    recebido = sum(p.total for p in pedidos if p.pago)
    telefone = cliente.telefone or next(
        (p.online.telefone for p in pedidos if p.online and p.online.telefone), ""
    )
    sugestao = mensagem = segmento = ""
    if pedidos and ranking:
        favorito = nomes[ranking[0]]
        segmento = (
            "Reativação: 30 dias ou mais"
            if dias >= 30
            else "Primeiro pedido" if len(pedidos) == 1 else "Cliente recorrente"
        )
        sugestao = f"Apresentar novamente {favorito}, o produto mais pedido."
        if len(top) > 1:
            sugestao += f" Alternativa já comprada: {nomes[top[1]]}."
        primeiro_nome = cliente.usuario.nome.strip().split()[0] if cliente.usuario.nome.strip() else ""
        mensagem = (
            f"Olá{(' ' + primeiro_nome) if primeiro_nome else ''}! "
            f"Que tal repetir seu pedido de {favorito}? "
            "Consulte nosso cardápio para conferir as opções de hoje."
        )

    linha = [
        cliente.id, cliente.usuario.nome, telefone, cliente.usuario.login,
        cliente.bairro, "Ativa" if cliente.usuario.ativo else "Desativada",
        len(pedidos), sum(p.status == "entregue" for p in pedidos),
        ultima, dias,
        reais(total) if pedidos else None,
        reais(recebido) if pedidos else None,
        (reais(total) / len(pedidos)).quantize(Decimal("0.01")) if pedidos else None,
        "\n".join(f"{nomes[k]} — {frequencia[k]} un." for k in top),
        segmento, sugestao, mensagem,
    ]
    detalhes = [
        [cliente.id, cliente.usuario.nome, posicao, nomes[k],
         frequencia[k], presencas[k], reais(valores[k])]
        for posicao, k in enumerate(ranking, 1)
    ]
    return linha, detalhes


def tabela(ws, titulo, nota, cabecalhos, linhas, larguras, nome,
           moedas=(), datas=(), identificadores=()):
    limite = get_column_letter(len(cabecalhos))
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.sheet_properties.tabColor = "A8C4D8"
    ws.merge_cells(f"A1:{limite}1")
    ws["A1"] = titulo
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color="334155")
    ws["A1"].fill = PatternFill("solid", fgColor="F1F5F9")
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 40
    ws.merge_cells(f"A2:{limite}2")
    ws["A2"] = nota
    ws["A2"].font = Font(name="Calibri", size=11, color="64748B")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 36
    ws.append([])
    ws.append(cabecalhos)
    for linha in linhas:
        ws.append(linha)

    # Texto do cadastro nunca é interpretado como fórmula do Excel.
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = texto(cell.value)
                cell.data_type = "s"
            if cell.row >= 5:
                cell.font = Font(
                    name="Calibri", size=11, color="475569",
                    bold=cell.column == 2,
                )
                cell.fill = PatternFill(
                    "solid", fgColor="FFFFFF" if cell.row % 2 else "F6F8FB"
                )
                cell.border = Border(bottom=Side(style="hair", color="E4EAF0"))
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.column in moedas:
                    cell.number_format = MOEDA
                    cell.alignment = Alignment(
                        horizontal="right", vertical="center", wrap_text=True
                    )
                elif cell.column in datas:
                    cell.number_format = "dd/mm/yyyy"
                elif cell.column in identificadores:
                    cell.number_format = "@"

    for cell in ws[4]:
        # Cores suaves separam cadastro, histórico e oportunidades.
        cor = (
            "E8EFF7" if nome != "TabelaLeads" or cell.column <= 6
            else "E9F2EE" if cell.column <= 13 else "F8F0E4"
        )
        cell.font = Font(name="Calibri", size=11, bold=True, color="334155")
        cell.fill = PatternFill("solid", fgColor=cor)
        cell.border = Border(bottom=Side(style="thin", color="CBD5E1"))
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 38
    for indice, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(indice)].width = largura
    for numero in range(5, ws.max_row + 1):
        # Ajuste ao conteúdo evita linhas altas para clientes sem histórico.
        maior = 1
        for cell in ws[numero]:
            if isinstance(cell.value, str):
                largura = max(8, int(larguras[cell.column - 1] - 3))
                linhas_texto = sum(
                    max(1, len(textwrap.wrap(parte, width=largura)))
                    for parte in cell.value.split("\n")
                )
                maior = max(maior, linhas_texto)
        ws.row_dimensions[numero].height = min(409, max(30, maior * 16 + 12))
    if linhas:
        table = Table(displayName=nome, ref=f"A4:{limite}{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight9", showRowStripes=True,
        )
        ws.add_table(table)
    else:
        ws.auto_filter.ref = f"A4:{limite}4"
    ws.freeze_panes = "C5"
    ws.print_title_rows = "1:4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:{limite}{max(ws.max_row, 4)}"


def gerar_planilha(clientes, pedidos, agora=None):
    agora = (agora or datetime.now(timezone.utc)).astimezone(BRASIL)
    por_cliente = defaultdict(list)
    for pedido in pedidos:
        if pedido.online:
            por_cliente[pedido.online.cliente_id].append(pedido)

    linhas, detalhes = [], []
    for cliente in clientes:
        linha, preferencias = resumir_cliente(
            cliente, por_cliente[cliente.id], agora.date()
        )
        linhas.append(linha)
        detalhes.extend(preferencias)

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads de clientes"
    emitido = agora.strftime("%d/%m/%Y %H:%M")
    tabela(
        ws, "Clientes e oportunidades de recompra",
        f"Exportado em {emitido} (Brasília). {len(clientes)} clientes. "
        "Histórico completo de pedidos online registrados, pagos ou não; "
        "valores dos pedidos incluem entrega. Sem pedidos: análise em branco.",
        ["ID cliente", "Nome", "Telefone", "Login", "Bairro", "Conta",
         "Pedidos registrados", "Pedidos entregues", "Último pedido",
         "Dias sem pedir", "Valor dos pedidos", "Valor marcado pago",
         "Ticket médio", "Top 3 produtos por quantidade", "Grupo sugerido",
         "Sugestão baseada no histórico", "Mensagem sugerida para revisão"],
        linhas,
        [12, 28, 21, 25, 22, 15, 16, 16, 17, 15, 20, 20, 18, 40, 29, 56, 65],
        "TabelaLeads", moedas=(11, 12, 13), datas=(9,), identificadores=(3, 4),
    )
    tabela(
        wb.create_sheet("Preferências"), "Produtos pedidos por cliente",
        "Fonte: itens dos pedidos online registrados. Ranking por unidades; "
        "empates por número de pedidos e nome. Nomes iguais são agrupados; "
        "nenhuma preferência é inferida para clientes sem pedidos. Sugestões "
        "não verificam estoque, descontos ou autorização de contato.",
        ["ID cliente", "Nome", "Posição", "Produto", "Unidades pedidas",
         "Pedidos com o produto", "Valor dos itens"],
        detalhes, [14, 30, 12, 45, 20, 24, 22],
        "TabelaPreferencias", moedas=(7,),
    )
    resultado = BytesIO()
    wb.save(resultado)
    resultado.seek(0)
    return resultado


@bp.get("/clientes/exportar.xlsx")
@role("admin")
def exportar_clientes():
    clientes = db.session.scalars(
        db.select(Cliente).options(joinedload(Cliente.usuario)).order_by(Cliente.id)
    ).all()
    pedidos = db.session.scalars(
        db.select(Pedido)
        .where(Pedido.online.has())
        .options(selectinload(Pedido.itens), joinedload(Pedido.online))
    ).all()
    arquivo = gerar_planilha(clientes, pedidos)
    resposta = send_file(
        arquivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"clientes_leads_{datetime.now(BRASIL):%Y-%m-%d_%H-%M}.xlsx",
        max_age=0,
    )
    resposta.headers["Cache-Control"] = "no-store, private"
    resposta.headers["X-Content-Type-Options"] = "nosniff"
    return resposta
