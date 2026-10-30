"""Relatórios administrativos por data de criação do pedido (Brasília)."""

from collections import defaultdict
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from zoneinfo import ZoneInfo

from flask import Blueprint, abort, render_template, request, send_file
from openpyxl import Workbook
from sqlalchemy.orm import joinedload, selectinload

from customers import role
from leads_export import tabela, reais
from models import Cliente, Pedido, Produto, db

bp = Blueprint("sales_reports", __name__)
BR = ZoneInfo("America/Sao_Paulo")
PERIODOS = {"dia": "Dia", "semana": "Semana", "mes": "Mês", "ano": "Ano"}
TIPOS = {"completo", "resumo", "categorias", "clientes", "produtos", "pedidos"}


def normalizar(nome):
    return " ".join(nome.split()).casefold()


def parametros():
    periodo = request.args.get("periodo", "dia")
    situacao = request.args.get("situacao", "pagos")
    if periodo not in PERIODOS or situacao not in {"pagos", "todos"}:
        abort(400, "Filtro inválido.")
    try:
        referencia = date.fromisoformat(
            request.args.get("referencia") or datetime.now(BR).date().isoformat()
        )
        if not 1900 <= referencia.year <= 9998:
            raise ValueError
        if periodo == "dia":
            inicio, fim = referencia, referencia + timedelta(days=1)
        elif periodo == "semana":
            inicio = referencia - timedelta(days=referencia.weekday())
            fim = inicio + timedelta(days=7)
        elif periodo == "mes":
            inicio = referencia.replace(day=1)
            fim = date(inicio.year + (inicio.month == 12), inicio.month % 12 + 1, 1)
        else:
            inicio, fim = date(referencia.year, 1, 1), date(referencia.year + 1, 1, 1)
    except (ValueError, OverflowError):
        abort(400, "Informe uma data válida.")
    return dict(periodo=periodo, situacao=situacao, referencia=referencia,
                inicio=inicio, fim=fim)


def limite_utc(dia):
    return datetime.combine(dia, time.min, BR).astimezone(timezone.utc).replace(tzinfo=None)


def local(criado):
    if criado.tzinfo is None:
        criado = criado.replace(tzinfo=timezone.utc)
    return criado.astimezone(BR)


def montar_relatorio(f):
    consulta = db.select(Pedido).where(
        Pedido.criado >= limite_utc(f["inicio"]),
        Pedido.criado < limite_utc(f["fim"]),
    ).options(selectinload(Pedido.itens), joinedload(Pedido.online))
    if f["situacao"] == "pagos":
        consulta = consulta.where(Pedido.pago.is_(True))
    pedidos = db.session.scalars(consulta.order_by(Pedido.criado, Pedido.id)).all()

    categorias_por_nome = defaultdict(set)
    for produto in db.session.scalars(
        db.select(Produto).options(joinedload(Produto.categoria))
    ):
        categorias_por_nome[normalizar(produto.nome)].add(produto.categoria.nome)
    clientes_ids = {p.online.cliente_id for p in pedidos if p.online}
    clientes_db = {
        c.id: c for c in db.session.scalars(
            db.select(Cliente).where(Cliente.id.in_(clientes_ids))
            .options(joinedload(Cliente.usuario))
        )
    } if clientes_ids else {}

    def grupo(nome):
        return dict(nome=nome, unidades=0, valor=0, ids=set())

    categorias, produtos, clientes = {}, {}, {}
    buckets = {}
    if f["periodo"] == "dia":
        for hora in range(24):
            buckets[hora] = dict(rotulo=f"{hora:02d}:00", pedidos=0, valor=0)
    elif f["periodo"] == "ano":
        for mes in range(1, 13):
            buckets[mes] = dict(rotulo=f"{mes:02d}/{f['inicio'].year}", pedidos=0, valor=0)
    else:
        dia = f["inicio"]
        while dia < f["fim"]:
            buckets[dia] = dict(rotulo=dia.strftime("%d/%m/%Y"), pedidos=0, valor=0)
            dia += timedelta(days=1)

    detalhe = []
    total = subtotal = taxas = pago = unidades = salao = sem_categoria = 0
    for p in pedidos:
        criado = local(p.criado)
        chave = criado.hour if f["periodo"] == "dia" else criado.month if f["periodo"] == "ano" else criado.date()
        buckets[chave]["pedidos"] += 1
        buckets[chave]["valor"] += p.total
        total += p.total
        subtotal += p.subtotal
        taxas += p.total - p.subtotal
        pago += p.total if p.pago else 0
        if not p.online:
            salao += p.total
        for item in p.itens:
            chave_nome = normalizar(item.nome)
            possiveis = categorias_por_nome.get(chave_nome, set())
            categoria = next(iter(possiveis)) if len(possiveis) == 1 else "Categoria não identificada"
            valor = item.quantidade * item.preco
            if len(possiveis) != 1:
                sem_categoria += valor
            g = categorias.setdefault(categoria, grupo(categoria))
            prod = produtos.setdefault(chave_nome, {**grupo(item.nome), "categorias": set()})
            prod["categorias"].add(categoria)
            for agregado in (g, prod):
                agregado["unidades"] += item.quantidade
                agregado["valor"] += valor
                agregado["ids"].add(p.id)
            unidades += item.quantidade
        nome_cliente = ""
        if p.online:
            cid = p.online.cliente_id
            cadastro = clientes_db.get(cid)
            nome_cliente = cadastro.usuario.nome if cadastro else p.online.destinatario
            c = clientes.setdefault(cid, dict(
                id=cid, nome=nome_cliente,
                telefone=(cadastro.telefone if cadastro else "") or p.online.telefone,
                pedidos=0, valor=0, ultimo=criado.date(),
            ))
            c["pedidos"] += 1
            c["valor"] += p.total
            c["ultimo"] = max(c["ultimo"], criado.date())
        detalhe.append(dict(
            id=p.id, criado=criado.replace(tzinfo=None),
            origem={"entrega": "Tele-entrega", "retirada": "Retirada"}.get(
                p.online.modalidade if p.online else "", "Salão"),
            comanda=p.mesa if not p.online else "", cliente=nome_cliente,
            status=p.status, pago=p.pago, subtotal=p.subtotal,
            taxa=p.total-p.subtotal, total=p.total,
        ))

    categorias = sorted(categorias.values(), key=lambda x: (-x["valor"], x["nome"].casefold()))
    produtos = sorted(produtos.values(), key=lambda x: (-x["unidades"], -x["valor"], x["nome"].casefold()))
    clientes = sorted(clientes.values(), key=lambda x: (-x["valor"], -x["pedidos"], x["id"]))
    return dict(
        filtros=f, pedidos=detalhe, categorias=categorias, produtos=produtos,
        clientes=clientes, evolucao=list(buckets.values()), quantidade=len(pedidos),
        total=total, subtotal=subtotal, taxas=taxas, recebido=pago, a_receber=total-pago,
        ticket=int((Decimal(total) / len(pedidos)).quantize(
            Decimal("1"), rounding=ROUND_HALF_UP
        )) if pedidos else 0,
        unidades=unidades, salao=salao, sem_categoria=sem_categoria,
    )


@bp.get("/relatorios-vendas")
@role("admin")
def pagina():
    f = parametros()
    return render_template("relatorios_vendas.html", r=montar_relatorio(f),
                           periodos=PERIODOS, fim_inclusivo=f["fim"]-timedelta(days=1))


def excel(r, tipo):
    wb = Workbook()
    wb.remove(wb.active)
    f = r["filtros"]
    nota = (
        f"{f['inicio']:%d/%m/%Y} a {f['fim']-timedelta(days=1):%d/%m/%Y} · "
        f"{'Somente pagos' if f['situacao']=='pagos' else 'Todos os pedidos'} · "
        "Data de criação do pedido, Brasília. Não é relatório de entradas por data de pagamento."
    )
    def aba(nome, titulo, cabecalhos, linhas, larguras, moedas=(), datas=(), ids=()):
        tabela(wb.create_sheet(nome), titulo, nota, cabecalhos, linhas, larguras,
               "Relatorio"+str(len(wb.worksheets)), moedas=moedas,
               datas=datas, identificadores=ids)
    if tipo in {"completo", "resumo"}:
        aba("Resumo", "Resumo das vendas", ["Indicador", "Valor"], [
            ["Pedidos", r["quantidade"]], ["Unidades", r["unidades"]],
            ["Produtos (sem entrega)", reais(r["subtotal"])],
            ["Taxas de entrega", reais(r["taxas"])],
            ["Total dos pedidos", reais(r["total"])],
            ["Valor marcado como pago", reais(r["recebido"])],
            ["Valor não pago", reais(r["a_receber"])],
            ["Ticket médio", reais(r["ticket"])],
        ], [36, 25])
        for row in range(7, 13):
            wb["Resumo"].cell(row, 2).number_format = '"R$" #,##0.00'
        aba("Evolução", "Vendas ao longo do período", ["Período", "Pedidos", "Valor"],
            [[x["rotulo"], x["pedidos"], reais(x["valor"])] for x in r["evolucao"]],
            [22, 18, 24], moedas=(3,), ids=(1,))
    if tipo in {"completo", "categorias"}:
        aba("Categorias", "Vendas por categoria atual", ["Categoria", "Unidades", "Pedidos com itens", "Valor dos produtos"],
            [[x["nome"], x["unidades"], len(x["ids"]), reais(x["valor"])] for x in r["categorias"]],
            [42, 18, 24, 26], moedas=(4,))
        wb["Categorias"]["A2"] = nota + " Categorias associadas pelo nome ao cadastro atual; renomeações e categorias históricas não são recuperáveis. Entrega excluída dos valores."
        wb["Categorias"].row_dimensions[2].height = 64
    if tipo in {"completo", "produtos"}:
        aba("Produtos", "Produtos mais vendidos por unidades", ["Produto", "Categoria atual", "Unidades", "Pedidos com produto", "Valor dos itens"],
            [[x["nome"], ", ".join(sorted(x["categorias"])), x["unidades"], len(x["ids"]), reais(x["valor"])] for x in r["produtos"]],
            [40, 35, 18, 24, 24], moedas=(5,))
        wb["Produtos"]["A2"] = nota + " Nomes iguais são agrupados; produtos renomeados podem ficar separados. Categoria baseada no cadastro atual. Entrega excluída."
        wb["Produtos"].row_dimensions[2].height = 58
    if tipo in {"completo", "clientes"}:
        aba("Clientes", "Melhores clientes por valor de pedidos", ["ID cliente", "Nome", "Telefone", "Pedidos", "Valor dos pedidos", "Ticket médio", "Último pedido"],
            [[x["id"], x["nome"], x["telefone"], x["pedidos"], reais(x["valor"]),
              reais(x["valor"])/x["pedidos"], x["ultimo"]] for x in r["clientes"]],
            [15, 32, 22, 16, 25, 22, 20], moedas=(5, 6), datas=(7,), ids=(3,))
        wb["Clientes"]["A2"] = nota + " Somente clientes identificados de pedidos online; pedidos do salão não são associados a um cliente."
        wb["Clientes"].row_dimensions[2].height = 52
    if tipo in {"completo", "pedidos"}:
        aba("Pedidos", "Detalhamento de pedidos", ["Pedido", "Criado em Brasília", "Origem", "Comanda", "Cliente", "Situação", "Pagamento", "Produtos", "Entrega", "Total"],
            [[x["id"], x["criado"], x["origem"], x["comanda"], x["cliente"], x["status"],
              "Pago" if x["pago"] else "Não pago", reais(x["subtotal"]), reais(x["taxa"]), reais(x["total"])] for x in r["pedidos"]],
            [14, 25, 20, 18, 32, 22, 17, 23, 20, 23], moedas=(8, 9, 10), datas=(2,), ids=(4,))
        for row in range(5, wb["Pedidos"].max_row+1):
            wb["Pedidos"].cell(row, 2).number_format = "dd/mm/yyyy hh:mm"
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


@bp.get("/relatorios-vendas/exportar.xlsx")
@role("admin")
def exportar():
    tipo = request.args.get("tipo", "completo")
    if tipo not in TIPOS:
        abort(400, "Relatório inválido.")
    f = parametros()
    resposta = send_file(
        excel(montar_relatorio(f), tipo),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"vendas_{tipo}_{f['periodo']}_{f['inicio']:%Y-%m-%d}.xlsx",
        max_age=0,
    )
    resposta.headers["Cache-Control"] = "no-store, private"
    resposta.headers["X-Content-Type-Options"] = "nosniff"
    return resposta
